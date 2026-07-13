from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from build_deang_sour_tea_prd import (
    BLUE,
    DARK_BLUE,
    INK,
    MUTED,
    add_bullet,
    add_callout,
    add_number,
    add_para,
    add_table,
    configure_doc,
    paragraph_border_bottom,
    set_run_font,
)


OUT = Path("德昂族酸茶数字化互动体验平台_技术方案.docx")


def add_header_footer(doc):
    sec = doc.sections[0]
    hp = sec.header.paragraphs[0]
    hp.text = ""
    r = hp.add_run("技术方案 | 德昂族酸茶数字化互动体验平台")
    set_run_font(r, size=9, color=MUTED)
    paragraph_border_bottom(hp, color="D8E0EA", size="4", space="4")

    fp = sec.footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = fp.add_run("Technical Solution v1.0")
    set_run_font(r, size=9, color=MUTED)


def add_masthead(doc):
    add_para(doc, "技术方案", size=23, bold=True, after=4)
    add_para(doc, "德昂族酸茶数字化互动体验平台", size=15, color=MUTED, after=14)
    rows = [
        ("文档版本", "v1.0"),
        ("对应PRD", "德昂族酸茶非遗科普、互动体验与电商转化平台 PRD v2.1"),
        ("技术路线", "Vue 3 + Vite + Java 17 + Spring Boot 3 + MyBatis-Plus + MySQL 8"),
        ("核心身份", "普通用户、商家、管理员"),
        ("编写日期", date.today().strftime("%Y-%m-%d")),
        ("适用阶段", "毕业设计概要设计、详细设计、开发实施与答辩说明"),
    ]
    for label, value in rows:
        p = add_para(doc, after=2)
        r1 = p.add_run(f"{label}: ")
        set_run_font(r1, size=11, bold=True)
        r2 = p.add_run(value)
        set_run_font(r2, size=11)
    rule = add_para(doc, after=12)
    paragraph_border_bottom(rule)


def build():
    doc = Document()
    configure_doc(doc)
    add_header_footer(doc)
    add_masthead(doc)

    doc.add_heading("1. 技术目标", level=1)
    add_callout(
        doc,
        "总体目标",
        "构建一个前后端分离的德昂族酸茶非遗平台，支持科普内容展示、H5互动体验、线下预约核销、酸茶电商交易、商家经营管理和管理员后台监管，满足毕业设计对页面数量、动态交互、数据库、权限管理、兼容测试和系统发布的要求。",
    )
    for item in [
        "前台用户侧：完成酸茶科普、H5游戏入口、商城浏览、购物车、订单、预约、评价等功能。",
        "商家侧：完成商家入驻、商品管理、库存价格、订单发货、售后处理、经营数据查看。",
        "管理员侧：完成用户管理、商家审核、商品审核、内容管理、订单监管、数据看板和权限控制。",
        "后端服务：提供稳定的RESTful API、JWT登录鉴权、RBAC权限控制、MySQL数据持久化和模拟支付订单流转。",
    ]:
        add_bullet(doc, item)

    doc.add_heading("2. 总体架构", level=1)
    add_table(
        doc,
        ["层级", "组成", "职责"],
        [
            ("表现层", "PC Web、移动端Web、H5互动页、后台管理页", "展示内容、承载交互、提交表单、发起接口请求"),
            ("前端应用层", "Vue 3、Vite、Vue Router、Pinia、Element Plus、ECharts", "路由管理、状态管理、页面组件、表格表单、数据可视化"),
            ("接口服务层", "Spring Boot 3、Spring MVC、RESTful API", "处理业务请求、参数校验、统一响应、异常处理"),
            ("业务服务层", "用户、内容、H5、预约核销、商品、购物车、订单、售后、数据统计", "封装核心业务规则与事务逻辑"),
            ("权限安全层", "Spring Security、JWT、RBAC", "登录认证、角色鉴权、接口访问控制、密码加密"),
            ("数据访问层", "MyBatis-Plus、MySQL 8", "数据表映射、CRUD、分页查询、事务持久化"),
            ("基础设施层", "Nginx、Spring Boot Jar、文件上传目录、服务器", "静态资源部署、后端运行、图片文件存储、系统访问"),
        ],
        [1300, 3400, 4660],
    )
    add_para(
        doc,
        "系统采用前后端分离架构。前端通过Axios调用后端RESTful接口，后端以Spring Boot提供业务服务，MySQL保存结构化数据。H5互动体验可以作为前端独立路由或小程序/H5独立入口接入，同一套后端用户、海报、预约、订单能力可复用。"
    )

    doc.add_heading("3. 技术选型", level=1)
    add_table(
        doc,
        ["类别", "技术", "选用理由"],
        [
            ("前端框架", "Vue 3 + Vite", "开发效率高，适合组件化页面、H5和后台管理系统"),
            ("UI组件", "Element Plus", "表格、表单、弹窗、分页成熟，适合管理员和商家后台"),
            ("状态管理", "Pinia", "管理登录用户、角色、购物车、订单草稿等前端状态"),
            ("路由", "Vue Router", "区分前台、用户中心、商家中心、管理员后台路由"),
            ("请求库", "Axios", "统一封装Token、错误处理和接口请求"),
            ("图表", "ECharts", "实现用户、订单、预约、核销、销量等数据看板"),
            ("后端框架", "Java 17 + Spring Boot 3", "生态成熟，适合毕业设计完整业务系统"),
            ("权限安全", "Spring Security + JWT", "实现登录认证、Token校验与接口权限控制"),
            ("ORM", "MyBatis-Plus", "简化CRUD、分页和条件查询，便于快速开发"),
            ("数据库", "MySQL 8", "适合用户、商品、订单、预约、内容等关系型数据"),
            ("文件上传", "本地文件存储", "毕设阶段实现简单，可扩展至OSS/COS"),
            ("支付", "模拟支付", "避免真实支付资质复杂度，完整演示订单状态流转"),
            ("部署", "Nginx + Spring Boot Jar", "前端静态部署，后端Jar包运行，结构清晰"),
        ],
        [1500, 2600, 5260],
        font_size=9,
    )

    doc.add_heading("4. 前端方案", level=1)
    add_table(
        doc,
        ["前端区域", "主要页面", "关键实现"],
        [
            ("前台展示", "首页、酸茶故事、制作技艺、传承人/工坊、用户作品、关于项目", "响应式布局、栏目导航、内容详情、图片展示、评论入口"),
            ("H5互动", "加载页、源起之魂、自然之魂、技艺之魂、火塘终章、海报生成", "移动端竖屏适配、CSS动画/Canvas、拖拽/长按/选择题、海报保存"),
            ("电商前台", "商城首页、商品列表、商品详情、购物车、确认订单、模拟支付、我的订单", "商品筛选、购物车状态、订单提交、支付状态切换"),
            ("用户中心", "我的预约、我的配方、我的海报、我的订单、售后、评价", "登录态校验、个人数据接口、订单状态标签"),
            ("商家中心", "商品管理、库存管理、订单处理、发货、售后、经营数据", "Element Plus表格表单、上传图片、分页查询、状态操作"),
            ("管理员后台", "用户管理、商家审核、商品审核、内容管理、订单监管、数据看板", "RBAC菜单、审核流、统计图表、批量操作"),
        ],
        [1500, 3300, 4560],
        font_size=9,
    )
    add_table(
        doc,
        ["前端目录", "说明"],
        [
            ("src/api", "按业务模块封装接口，如auth、product、order、booking、content"),
            ("src/router", "配置前台、用户、商家、管理员路由与路由守卫"),
            ("src/stores", "Pinia状态，如userStore、cartStore、permissionStore"),
            ("src/views/site", "科普展示与内容页面"),
            ("src/views/h5", "H5互动体验页面"),
            ("src/views/shop", "商城、商品、购物车、订单页面"),
            ("src/views/merchant", "商家中心页面"),
            ("src/views/admin", "管理员后台页面"),
            ("src/components", "通用组件，如商品卡片、上传组件、状态标签、数据卡片"),
        ],
        [2200, 7160],
    )

    doc.add_heading("5. 后端方案", level=1)
    add_table(
        doc,
        ["后端分层", "包/模块", "职责"],
        [
            ("Controller", "controller", "接收HTTP请求、参数校验、返回统一响应"),
            ("Service", "service", "封装业务逻辑、事务处理、权限内业务校验"),
            ("Mapper", "mapper", "MyBatis-Plus数据访问接口"),
            ("Entity", "entity", "数据库实体对象"),
            ("DTO/VO", "dto、vo", "请求入参与响应出参，避免直接暴露实体"),
            ("Security", "security", "JWT过滤器、用户认证、权限校验、密码加密"),
            ("Common", "common", "统一响应、异常处理、分页对象、枚举、工具类"),
        ],
        [1500, 2600, 5260],
    )
    add_table(
        doc,
        ["业务模块", "后端能力"],
        [
            ("认证与用户", "注册、登录、密码加密、JWT签发、用户信息、角色绑定"),
            ("内容管理", "酸茶文章、分类、图片、首页推荐、评论留言、分页展示"),
            ("H5互动", "游戏进度、用户选择、配方生成、海报编号、行为埋点"),
            ("预约核销", "场次管理、预约提交、核销码生成、到店核销、评价反馈"),
            ("商品商城", "商品分类、商品列表、详情、搜索、上下架、库存价格"),
            ("购物车", "加入购物车、修改数量、删除、勾选、结算"),
            ("订单交易", "创建订单、模拟支付、取消、发货、收货、完成、退款/售后"),
            ("商家中心", "入驻申请、商品管理、订单处理、售后处理、经营统计"),
            ("管理员后台", "用户管理、商家审核、商品审核、订单监管、内容管理、数据看板"),
            ("统计看板", "用户数、订单额、商品销量、预约量、核销率、H5完成率"),
        ],
        [1600, 7760],
    )

    doc.add_heading("6. 身份权限设计", level=1)
    add_table(
        doc,
        ["角色", "角色编码", "主要权限"],
        [
            ("普通用户", "USER", "浏览内容、H5互动、生成海报、预约体验、购买商品、查看订单、申请售后、评价"),
            ("商家", "MERCHANT", "店铺资料、商品管理、库存维护、订单处理、发货、售后处理、经营数据"),
            ("管理员", "ADMIN", "用户管理、商家审核、商品审核、内容管理、订单监管、投诉处理、平台数据看板"),
        ],
        [1500, 1500, 6360],
    )
    for step in [
        "用户登录后，后端校验账号密码并签发JWT，Token中包含userId与role。",
        "前端保存Token，请求接口时通过Authorization请求头携带Token。",
        "Spring Security过滤器解析Token并写入认证上下文。",
        "后端通过注解或配置限制接口角色，例如商品审核仅ADMIN可访问，发货仅MERCHANT可访问。",
        "前端根据角色动态加载菜单，普通用户、商家、管理员进入不同工作台。",
    ]:
        add_number(doc, step)

    doc.add_heading("7. 数据库设计", level=1)
    add_table(
        doc,
        ["数据表", "关键字段", "说明"],
        [
            ("user", "id、username、password、phone、role、status、created_at", "用户与三类角色基础信息"),
            ("merchant", "id、user_id、shop_name、contact、license_url、audit_status", "商家入驻与审核信息"),
            ("content", "id、title、category_id、cover_url、body、status、created_at", "酸茶科普内容"),
            ("game_session", "id、user_id、progress、choices、completion_time、created_at", "H5互动进度与选择"),
            ("tea_profile", "id、user_id、recipe_type、keyword、guardian_no、poster_code", "专属茶配方与守护人编号"),
            ("booking", "id、user_id、session_date、people_count、status、verify_code", "线下体验预约与核销"),
            ("product", "id、merchant_id、category_id、name、price、stock、status", "商品信息"),
            ("cart_item", "id、user_id、product_id、quantity、checked", "购物车"),
            ("order", "id、order_no、user_id、merchant_id、total_amount、status、pay_status", "订单主表"),
            ("order_item", "id、order_id、product_id、product_name、price、quantity", "订单明细"),
            ("after_sale", "id、order_id、user_id、merchant_id、type、reason、status", "退款与售后"),
            ("feedback", "id、user_id、target_type、target_id、score、content", "商品、体验或内容评价"),
        ],
        [1550, 4400, 3410],
        font_size=8.5,
    )

    doc.add_heading("8. 核心接口规划", level=1)
    add_table(
        doc,
        ["模块", "接口示例", "说明"],
        [
            ("认证", "POST /api/auth/register；POST /api/auth/login", "注册、登录、获取Token"),
            ("内容", "GET /api/contents；GET /api/contents/{id}", "酸茶内容列表与详情"),
            ("H5", "POST /api/game/sessions；POST /api/game/profile", "保存游戏进度、生成茶配方"),
            ("预约", "POST /api/bookings；POST /api/bookings/{id}/verify", "创建预约、到店核销"),
            ("商品", "GET /api/products；GET /api/products/{id}", "商品列表与详情"),
            ("购物车", "POST /api/cart/items；PUT /api/cart/items/{id}", "加入购物车、修改数量"),
            ("订单", "POST /api/orders；POST /api/orders/{id}/pay/mock", "创建订单、模拟支付"),
            ("商家", "POST /api/merchant/apply；POST /api/merchant/products", "商家入驻、发布商品"),
            ("管理员", "PUT /api/admin/merchants/{id}/audit；PUT /api/admin/products/{id}/audit", "商家审核、商品审核"),
            ("统计", "GET /api/dashboard/admin；GET /api/dashboard/merchant", "平台与商家数据看板"),
        ],
        [1500, 3400, 4460],
        font_size=8.8,
    )

    doc.add_heading("9. 关键业务流程", level=1)
    add_table(
        doc,
        ["流程", "步骤"],
        [
            ("普通用户购买", "浏览商品→加入购物车→确认订单→模拟支付→商家发货→用户确认收货→评价"),
            ("商家发布商品", "商家登录→填写商品资料→上传图片→提交审核→管理员审核→上架销售"),
            ("管理员审核", "查看待审核商家/商品→核对资料→通过或驳回→记录审核结果"),
            ("H5到预约", "完成H5→生成配方和海报→点击预约→选择场次→生成核销码→到店核销"),
            ("售后处理", "用户提交售后→商家处理→管理员必要时介入→更新订单与售后状态"),
        ],
        [1800, 7560],
    )

    doc.add_heading("10. 安全与异常处理", level=1)
    for item in [
        "密码使用BCrypt加密保存，数据库不保存明文密码。",
        "所有需要登录的接口必须校验JWT，Token过期后要求重新登录。",
        "通过RBAC限制接口访问，商家只能操作自己店铺的商品和订单。",
        "订单金额以后端商品价格为准计算，前端传入金额不作为最终可信数据。",
        "上传文件限制类型和大小，商品图、内容图分别存储并记录路径。",
        "统一异常处理，返回标准错误码和提示信息，便于前端展示。",
        "模拟支付仅用于毕设演示，不接入真实支付渠道和资金清算。",
    ]:
        add_bullet(doc, item)

    doc.add_heading("11. 部署方案", level=1)
    add_table(
        doc,
        ["组件", "部署方式"],
        [
            ("前端", "Vite打包生成dist目录，通过Nginx部署静态资源"),
            ("后端", "Spring Boot打包为Jar，使用java -jar或systemd运行"),
            ("数据库", "MySQL 8部署在本机或云服务器，初始化表结构和测试数据"),
            ("文件", "上传文件存储在服务器本地uploads目录，数据库保存相对路径"),
            ("反向代理", "Nginx将/api请求代理到Spring Boot服务，其他请求返回前端页面"),
        ],
        [1800, 7560],
    )

    doc.add_heading("12. 测试方案", level=1)
    add_table(
        doc,
        ["测试类型", "重点"],
        [
            ("功能测试", "注册登录、H5配方生成、预约核销、商品购买、订单发货、售后、审核"),
            ("权限测试", "普通用户、商家、管理员访问边界，验证越权接口不可用"),
            ("兼容测试", "Chrome、Edge、Safari/Firefox；微信内置浏览器移动端H5"),
            ("数据测试", "分页、搜索、筛选、库存扣减、订单状态流转、统计接口"),
            ("异常测试", "Token过期、库存不足、重复支付、非法订单、审核驳回、文件上传失败"),
            ("演示测试", "准备种子账号、商品、订单、预约、核销码和数据看板样例"),
        ],
        [1800, 7560],
    )

    doc.add_heading("13. 开发里程碑", level=1)
    add_table(
        doc,
        ["阶段", "产出"],
        [
            ("第1阶段：基础框架", "Vue项目、Spring Boot项目、MySQL表结构、登录鉴权、角色路由"),
            ("第2阶段：内容与H5", "酸茶内容、H5流程、配方生成、海报编号、行为埋点"),
            ("第3阶段：预约核销", "预约场次、核销码、到店核销、评价反馈"),
            ("第4阶段：电商系统", "商品、购物车、订单、模拟支付、发货、售后"),
            ("第5阶段：商家与后台", "商家入驻、商品审核、订单监管、数据看板"),
            ("第6阶段：测试发布", "兼容测试、部署、演示账号、运行说明、答辩材料"),
        ],
        [2300, 7060],
    )

    doc.add_heading("14. 技术风险与对策", level=1)
    add_table(
        doc,
        ["风险", "对策"],
        [
            ("电商流程过大", "先实现商品、购物车、订单、模拟支付、发货五个核心闭环，优惠券、分账等作为后续扩展"),
            ("权限越权", "后端接口必须按角色和数据归属校验，商家操作订单时校验merchant_id"),
            ("H5动画耗时", "优先实现关键交互，复杂动画用CSS/Canvas简化，保证流程可演示"),
            ("数据库关系复杂", "先完成核心表，使用订单主表+明细表结构，售后表单独维护"),
            ("部署环境不稳定", "准备本地演示环境和云端部署两套方案，提前导入测试数据"),
        ],
        [2300, 7060],
    )

    doc.save(OUT)


if __name__ == "__main__":
    build()
    print(OUT.resolve())
