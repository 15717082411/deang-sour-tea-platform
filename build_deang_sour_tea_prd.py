from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("德昂族酸茶数字化互动体验平台_PRD.docx")

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(20, 20, 20)
MUTED = RGBColor(89, 89, 89)
LIGHT_FILL = "F2F4F7"
CALLOUT_FILL = "F4F6F9"


def set_run_font(run, size=None, color=None, bold=None, italic=None, east_asia="Microsoft YaHei"):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, val in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def paragraph_border_bottom(paragraph, color="AFC0D3", size="8", space="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)


def add_para(doc, text="", size=11, color=INK, bold=False, italic=False, after=6, before=0, align=None, style=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    if align is not None:
        p.alignment = align
    if text:
        run = p.add_run(text)
        set_run_font(run, size=size, color=color, bold=bold, italic=italic)
    return p


def set_cell_text(cell, text, bold=False, color=INK, size=10, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.1
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold)


def add_table(doc, headers, rows, widths, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT_FILL)
        set_cell_text(cell, header, bold=True, color=DARK_BLUE, size=font_size, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row_data in rows:
        row = table.add_row()
        for idx, value in enumerate(row_data):
            text = str(value)
            align = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 or len(text) <= 8 else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(row.cells[idx], text, size=font_size, align=align)
    add_para(doc, "", after=4)
    return table


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    run = p.add_run(text)
    set_run_font(run, size=11, color=INK)


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    run = p.add_run(text)
    set_run_font(run, size=11, color=INK)


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT_FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=11, bold=True, color=DARK_BLUE)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.1
    r2 = p2.add_run(body)
    set_run_font(r2, size=10.5, color=INK)
    add_para(doc, "", after=6)


def configure_doc(doc):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.1


def add_header_footer(doc):
    sec = doc.sections[0]
    hp = sec.header.paragraphs[0]
    hp.text = ""
    r = hp.add_run("产品需求文档 | 德昂族酸茶数字化互动体验平台")
    set_run_font(r, size=9, color=MUTED)
    paragraph_border_bottom(hp, color="D8E0EA", size="4", space="4")
    fp = sec.footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = fp.add_run("PRD v2.1")
    set_run_font(r, size=9, color=MUTED)


def add_masthead(doc):
    add_para(doc, "产品需求文档", size=23, bold=True, after=4)
    add_para(doc, "德昂族酸茶数字化互动体验平台", size=15, color=MUTED, after=14)
    rows = [
        ("文档版本", "v2.1，新增酸茶电商系统与三类身份权限"),
        ("核心对象", "德昂族酸茶制作技艺；史诗叙事仅作为酸茶体验的引导与仪式层"),
        ("产品形态", "PC Web展示站 + H5/微信小程序轻游戏 + 移动端预约核销 + 酸茶电商系统 + 后台管理"),
        ("参考资料", "《2027届毕业设计选题方向》与《3.16版本计划书》"),
        ("编写日期", date.today().strftime("%Y-%m-%d")),
        ("适用场景", "毕业设计立项、原型设计、开发排期与答辩评审"),
    ]
    for label, value in rows:
        p = add_para(doc, after=2)
        r1 = p.add_run(f"{label}: ")
        set_run_font(r1, size=11, bold=True)
        r2 = p.add_run(value)
        set_run_font(r2, size=11)
    rule = add_para(doc, after=12)
    paragraph_border_bottom(rule)


def build():
    doc = Document()
    configure_doc(doc)
    add_header_footer(doc)
    add_masthead(doc)

    doc.add_heading("1. 产品定位", level=1)
    add_callout(
        doc,
        "一句话定义",
        "以德昂族酸茶制作技艺为核心的数字化科普、互动体验与电商转化平台，通过线上轻游戏完成认知铺垫、专属茶配方生成与预约引流，通过线下酸茶工坊完成杀青、揉捻、发酵、品饮与纪念核销，并通过电商系统承接酸茶产品、文创周边与体验套餐的购买需求。",
    )
    add_para(
        doc,
        "本次重构不再聚焦泛茶文化或完整茶文化史，而是聚焦“德昂族酸茶”这一具体非遗技艺。平台的所有内容、交互和指标均围绕酸茶的认知、制作、体验、购买、转化与传承展开；《达古达楞格莱标》等史诗元素只承担叙事引导、仪式感与文化出处说明，不扩展为独立的大型史诗产品。"
    )

    doc.add_heading("1.1 核心问题", level=2)
    for item in [
        "德昂族酸茶认知度低：计划书调研显示，大量受访者未听说过德昂族史诗或酸茶制作技艺，酸茶缺少清晰、可传播的公众符号。",
        "酸茶技艺传承门槛高：杀青、揉捻、发酵等工序依赖手工经验，制作周期长，年轻用户缺少低门槛接触与参与机会。",
        "非遗体验冷启动难：芒市本地文化点位分散，传统“观光+手作”模式难以形成主动分享和线上引流。",
        "酸茶消费承接不足：用户被内容或H5激发兴趣后，缺少直接购买酸茶产品、体验套餐和文创周边的交易入口。",
        "毕业设计交付要求完整：需满足PC与移动端页面、动态交互、数据库、兼容测试、发布、演示视频与海报等系统化要求。",
    ]:
        add_bullet(doc, item)

    doc.add_heading("1.2 产品目标", level=2)
    add_table(
        doc,
        ["目标", "说明", "验收指标"],
        [
            ("聚焦酸茶", "把内容主线收敛到德昂族酸茶的来源、工艺、口味、体验与传承", "全站一级内容中酸茶相关占比≥80%"),
            ("低门槛认知", "用3-5分钟H5轻游戏讲清酸茶的核心工艺和体验价值", "游戏完成率≥35%，酸茶工艺题正确率持续提升"),
            ("线上引流", "通过专属茶配方、守护人编号、海报和核销码促成线下预约", "海报生成率≥25%，预约点击率≥5%"),
            ("线下闭环", "让用户在线下完成酸茶工坊三步技艺挑战并兑换纪念物", "核销成功率≥90%，NPS>40"),
            ("电商转化", "承接酸茶产品、文创周边、体验套餐的浏览、下单、发货与售后流程", "普通用户可下单，商家可管理商品和订单，管理员可审核与监管"),
            ("毕设达标", "完成网站、小程序/H5、移动端、后台与动态功能", "15+页面、移动端5+页面、动态功能4+项"),
        ],
        [1500, 5000, 2860],
    )

    doc.add_heading("2. 用户与场景", level=1)
    add_table(
        doc,
        ["角色", "需求", "典型场景"],
        [
            ("年轻文化体验用户", "寻找小众、可分享、能亲手参与的旅行体验", "在小红书/抖音看到H5海报，生成配方并预约芒市体验"),
            ("普通用户/芒市游客", "希望了解酸茶、预约体验、购买酸茶和文创", "浏览商品、加入购物车、提交订单、查看物流或到店自提"),
            ("商家/工坊", "希望增加酸茶体验客流与线上销售，同时维护商品和订单", "在商家中心发布酸茶商品、处理订单、管理库存和售后"),
            ("管理员", "需要监管内容、商家、商品、订单、用户和数据", "审核商家入驻、处理违规商品、查看平台数据"),
            ("指导教师/评审", "需要看到主题聚焦、功能完整、技术可实现的毕业设计", "查看PC站、移动端、后台、演示视频与运行说明"),
        ],
        [1600, 3400, 4360],
    )

    doc.add_heading("2.1 核心用户旅程", level=2)
    for step in [
        "用户通过分享链接、二维码或短视频入口进入H5游戏。",
        "用户完成序章、源起、自然、酸茶工坊、火塘终章等轻量交互，重点学习杀青、揉捻、45天发酵等酸茶工艺知识。",
        "系统基于用户选择生成专属酸茶配方、守护人编号、核销码和分享海报。",
        "用户点击预约线下，选择芒市酸茶工坊场次，并获得到店提醒。",
        "用户到店核销，工作人员读取画像卡片，发放任务包，引导完成酸茶制作体验。",
        "用户领取编号一致的纪念币/茶品，分享体验，系统沉淀评价与复购线索。",
        "用户进入酸茶商城，购买专属配方关联商品、非遗酸茶礼盒、文创周边或线下体验套餐。",
        "商家在商家中心处理订单、维护库存与发货信息，管理员对商家和商品进行审核监管。",
    ]:
        add_number(doc, step)

    doc.add_heading("3. 产品范围", level=1)
    add_table(
        doc,
        ["模块", "本期包含", "明确不做"],
        [
            ("酸茶内容站", "酸茶来源、工艺、传承人、工坊、预约入口、作品展示", "泛中国茶史、普洱/绿茶/红茶等横向百科"),
            ("H5轻游戏", "三片茶魂收集、酸茶工艺小游戏、配方生成、海报分享、预约跳转", "大型RPG、复杂战斗、长篇史诗阅读"),
            ("预约核销", "场次选择、用户信息、核销码、到店核验、状态更新", "完整OTA平台和支付清结算体系"),
            ("电商系统", "商品浏览、商品详情、购物车、模拟支付、订单、发货、售后、商家中心", "真实支付牌照、复杂分账清算、跨境电商"),
            ("后台管理", "内容、预约、核销、评论、问卷、商品、商家、订单、数据看板", "复杂CRM和多门店财务系统"),
            ("线下体验辅助", "任务包说明、NPC话术、物料清单、核销联动", "重资产景区建设和高成本MR硬件依赖"),
        ],
        [1400, 4700, 3260],
    )

    doc.add_heading("4. 信息架构与页面清单", level=1)
    add_table(
        doc,
        ["端", "页面/功能", "说明"],
        [
            ("PC Web", "首页、酸茶故事、制作技艺、传承人/工坊、体验预约、酸茶商城、用户作品、关于项目", "用于毕业设计展示与完整站点结构"),
            ("H5/小程序", "加载页、序章、源起之魂、自然之魂、技艺之魂、火塘终章、海报生成、预约页", "用于传播、引流与轻交互"),
            ("移动端Web", "首页、技艺页、游戏入口、预约页、商城页、购物车、我的核销码", "至少5页移动适配"),
            ("用户中心", "登录注册、密码找回、预约记录、我的配方、我的海报、我的订单、评价反馈", "支撑动态交互与用户沉淀"),
            ("商家中心", "商家入驻、商品管理、库存管理、订单处理、售后管理、经营数据", "支撑商家角色独立运营"),
            ("后台", "内容管理、预约管理、核销管理、商家审核、商品审核、订单监管、留言审核、问卷统计、数据看板", "支撑动态网站与运营管理"),
        ],
        [1200, 3200, 4960],
    )
    add_table(
        doc,
        ["毕业设计页面要求", "本PRD对应实现"],
        [
            ("至少4个不同布局页面", "首页、酸茶技艺详情、H5游戏页、预约/后台页"),
            ("至少15个高质量页面", "PC 8页 + H5 8页 + 商城/用户/商家/后台若干页，可满足15+"),
            ("至少5个移动端页面", "移动首页、技艺页、游戏入口、预约页、商城页、购物车、我的订单"),
            ("包含一级导航页、栏目页、内容页", "PC Web信息架构完整覆盖"),
        ],
        [2600, 6760],
    )

    doc.add_heading("5. 核心功能需求", level=1)
    add_table(
        doc,
        ["编号", "功能", "需求描述", "优先级"],
        [
            ("F1", "酸茶主题首页", "展示德昂族酸茶定位、核心视觉、H5入口、预约入口、工艺亮点与工坊信息", "P0"),
            ("F2", "酸茶内容库", "提供酸茶来源、制作步骤、口味特征、传承现状、工坊地图等内容页", "P0"),
            ("F3", "H5序章与引导", "通过剪影、水墨纹理、音效和简短文案建立酸茶寻踪氛围", "P0"),
            ("F4", "源起之魂交互", "竹简排序或擦拭信笺，完成德昂族与酸茶来源的低门槛认知", "P1"),
            ("F5", "自然之魂交互", "古茶林探索、茶王树寻找、长按连接，用于建立酸茶原料与自然场景认知", "P1"),
            ("F6", "技艺之魂交互", "杀青控温QTE、揉捻手势、45天发酵选择，突出酸茶核心工艺", "P0"),
            ("F7", "配方生成", "基于用户选择生成德昂酸茶+小青柠/薄荷、玫瑰/桂圆、原味等专属配方", "P0"),
            ("F8", "海报与核销码", "生成含守护人编号、茶配方、核销码的分享海报，可保存与分享", "P0"),
            ("F9", "线下预约", "用户选择场次、提交联系方式、生成预约状态，并可跳转到核销页", "P0"),
            ("F10", "到店核销", "工作人员扫码读取用户配方、标签、终章关键词，并更新核销状态", "P0"),
            ("F11", "后台管理", "管理酸茶内容、预约、核销、评论、问卷、用户反馈与首页推荐", "P0"),
            ("F12", "数据看板", "展示PV/UV、游戏完成率、海报生成率、预约点击率、核销率、NPS等指标", "P1"),
            ("F13", "酸茶商城", "展示酸茶礼盒、体验套餐、文创周边，支持分类、搜索、筛选和商品详情", "P0"),
            ("F14", "购物车与下单", "普通用户可加入购物车、填写收货信息、提交订单并使用模拟支付完成购买", "P0"),
            ("F15", "订单中心", "普通用户可查看待支付、待发货、待收货、已完成、退款/售后订单", "P0"),
            ("F16", "商家中心", "商家可发布商品、维护库存价格、处理订单、录入物流、处理售后", "P0"),
            ("F17", "商家入驻审核", "管理员审核商家资质、启停商家账号，并监管商品上架状态", "P0"),
            ("F18", "电商数据统计", "统计商品销量、订单金额、转化率、商家销售排行和退款情况", "P1"),
        ],
        [700, 1550, 5910, 1200],
        font_size=8.8,
    )

    doc.add_heading("5.1 H5游戏流程", level=2)
    add_table(
        doc,
        ["阶段", "用户动作", "酸茶聚焦点", "关键数据"],
        [
            ("加载/序章", "进入链接、点击开始寻踪", "建立德昂酸茶的神秘入口，不展开泛茶文化", "渠道、设备、启动时间、开始率"),
            ("源起之魂", "擦拭信笺、竹简排序", "理解德昂族与茶的关系", "完成耗时、尝试次数、选项偏好"),
            ("自然之魂", "驱散迷雾、寻找茶王树", "感知古茶林与酸茶原料", "点击热区、长按成功率"),
            ("技艺之魂", "杀青、揉捻、发酵三步挑战", "学习酸茶核心工艺：火候、揉捻、45天发酵", "点击频率、手势进度、天数选择"),
            ("火塘终章", "合成碎片，回答“茶魂是什么”", "形成酸茶体验的个人意义", "完成率、关键词、情感倾向"),
            ("海报预约", "保存海报或点击预约", "把虚拟配方转为线下酸茶体验", "海报生成、保存、预约点击"),
        ],
        [1350, 2500, 3210, 2300],
        font_size=9,
    )

    doc.add_heading("5.2 线下酸茶体验联动", level=2)
    add_table(
        doc,
        ["站点", "体验内容", "系统联动"],
        [
            ("村口/工坊接待", "核验守护人海报，发放魂引牌、地图、任务包", "扫码读取核销码、配方与用户标签"),
            ("源起展示区", "竹简排序、史诗片段音频、酸茶文化来源说明", "记录到店进度，可展示线上答案"),
            ("古茶林/原料区", "识别茶树、闻香、触摸茶具与茶叶", "引导拍照分享，沉淀UGC"),
            ("酸茶坊", "杀青控温、揉捻、放入45天竹筒，体验酸茶制作", "与线上技艺之魂呼应，记录工艺完成"),
            ("火塘/品饮区", "写下关键词、盖茶祖印章、品饮酸茶、领取纪念币", "核销完成，触发评价、复购与分享激励"),
        ],
        [1600, 5200, 2560],
    )

    doc.add_heading("5.3 电商系统流程", level=2)
    add_table(
        doc,
        ["流程", "普通用户", "商家", "管理员"],
        [
            ("商品浏览", "浏览酸茶礼盒、专属配方商品、体验套餐、文创周边", "维护商品标题、图片、价格、库存、规格", "审核商品分类、敏感内容与上架状态"),
            ("购物车/下单", "加入购物车，填写地址，提交订单，模拟支付", "接收订单，确认库存，准备发货或到店核销", "查看异常订单，处理投诉"),
            ("订单履约", "查看订单状态、物流或自提核销信息", "录入物流单号，更新发货/完成状态", "监管订单流转与商家履约效率"),
            ("售后评价", "申请退款/售后，确认收货，发表评价", "处理售后，回复评价", "介入争议，管理评价违规内容"),
            ("数据复盘", "查看购买记录与推荐商品", "查看销量、库存、退款、转化数据", "查看平台交易、商家排行、商品转化数据"),
        ],
        [1500, 3100, 2500, 2260],
        font_size=8.8,
    )

    doc.add_heading("6. 数据与权限", level=1)
    add_table(
        doc,
        ["数据对象", "关键字段", "用途"],
        [
            ("User", "id、昵称、手机号、角色、来源渠道", "注册登录、预约、画像关联"),
            ("GameSession", "session_id、用户id、场景进度、完成时间、流失节点", "漏斗分析与体验优化"),
            ("TeaProfile", "用户标签、配方类型、终章关键词、守护人编号", "专属海报、到店接待话术"),
            ("PosterCode", "核销码、状态、生成时间、预约id", "线上线下闭环凭证"),
            ("Booking", "场次、人数、联系方式、状态、核销时间", "预约与现场运营"),
            ("Merchant", "商家id、店铺名、联系人、资质状态、启停状态", "商家入驻、审核与经营管理"),
            ("Product", "商品id、商家id、分类、标题、图片、价格、库存、状态", "商城展示、商品审核与库存管理"),
            ("Cart", "用户id、商品id、规格、数量、勾选状态", "购物车与下单前暂存"),
            ("Order", "订单id、用户id、商家id、金额、状态、地址、支付状态", "订单履约、售后与统计"),
            ("OrderItem", "订单id、商品id、规格、数量、单价", "订单明细与销量统计"),
            ("Content", "标题、分类、正文、图片、状态、发布时间", "酸茶内容动态更新"),
            ("Feedback", "评分、NPS、评论、复购意向", "服务质量与答辩数据支撑"),
        ],
        [1500, 4300, 3560],
    )
    add_table(
        doc,
        ["角色", "权限"],
        [
            ("普通用户", "浏览内容、玩H5、生成海报、预约体验、购买商品、查看订单、申请售后、提交评价"),
            ("商家", "提交入驻资料、管理店铺、发布商品、维护库存价格、处理订单、发货、处理售后、查看经营数据"),
            ("管理员", "管理用户、审核商家、审核商品、管理内容、监管订单、处理投诉、管理预约核销、查看平台数据"),
        ],
        [1800, 7560],
    )

    doc.add_heading("7. 非功能需求", level=1)
    for item in [
        "主题约束：所有页面和交互必须围绕德昂族酸茶，避免扩展成泛茶文化百科；涉及史诗内容时必须服务于酸茶体验理解。",
        "前端实现：PC站与后台以HTML5、CSS3、JavaScript为主；H5游戏需适配微信内置浏览器、iOS Safari和主流安卓浏览器。",
        "移动适配：H5游戏需竖屏优先，关键按钮触达区不小于44px，低端机加载时长需控制在可接受范围。",
        "电商边界：毕业设计阶段采用模拟支付或订单状态流转，不接入真实支付清算；订单、发货、退款流程需可演示。",
        "动态功能：至少实现注册登录/密码找回、预约核销、内容动态更新、商品管理、购物车、订单管理、留言评论、问卷反馈、分页展示中的4项。",
        "兼容发布：完成Chrome、Edge、Safari或Firefox三类主流浏览器兼容测试，并部署可访问版本。",
        "文化真实性：酸茶工艺、传承人信息、口味配方和线下话术需经过资料核对或传承人/指导教师审核。",
        "交易安全：普通用户、商家、管理员需严格隔离权限；商家不得访问其他商家订单，普通用户不得访问后台接口。",
        "隐私合规：收集手机号、微信头像、昵称、收货地址、订单和行为数据时需提示用途，后台仅展示运营必要信息。",
    ]:
        add_bullet(doc, item)

    doc.add_heading("8. 验收指标", level=1)
    add_table(
        doc,
        ["类别", "指标"],
        [
            ("产品完成度", "PC/H5/移动端/后台主要流程可演示，用户可从H5完成配方生成并进入预约。"),
            ("页面数量", "全站不少于15个高质量页面，移动端不少于5页，至少4个不同布局页面完成原型与效果图。"),
            ("酸茶聚焦", "首页、H5技艺之魂、内容页、预约页、后台数据均围绕德昂族酸茶展开。"),
            ("动态功能", "注册登录、预约、核销、内容管理、留言/评价、问卷/数据统计至少4项可运行。"),
            ("电商系统", "普通用户、商家、管理员三类身份可登录到对应功能；商品浏览、购物车、下单、模拟支付、发货、售后流程可演示。"),
            ("关键漏斗", "可统计游戏启动、技艺之魂完成、海报生成、预约点击、到店核销等事件。"),
            ("交易指标", "可统计商品浏览、加入购物车、下单、支付完成、退款售后、商家销售额等事件。"),
            ("运营指标", "目标线上至线下转化效率≥5%，NPS>40，复购及转介绍率≥20%。"),
            ("答辩材料", "提交运行说明、演示视频、兼容测试记录、3000px*2000px/300dpi/JPEG竖幅海报。"),
        ],
        [2000, 7360],
    )

    doc.add_heading("9. 里程碑", level=1)
    add_table(
        doc,
        ["阶段", "产出", "重点"],
        [
            ("需求重构", "酸茶聚焦PRD、信息架构、功能清单", "收敛范围，删除泛茶文化叙述"),
            ("原型设计", "PC 4类页面、H5完整流程、移动端5页、商城/商家中心/后台关键页", "保证页面数量、交易流程与交互闭环"),
            ("视觉设计", "德昂酸茶视觉体系、海报模板、H5场景效果图", "突出酸茶、茶林、工坊、火塘等真实场景"),
            ("前端开发", "PC站、H5游戏、移动适配、商城、购物车、订单页、商家中心", "完成技艺之魂、海报生成与电商购买关键体验"),
            ("后端开发", "用户、内容、预约、核销、商品、购物车、订单、反馈、数据看板", "支撑动态功能、交易流程和运营闭环"),
            ("测试发布", "三浏览器兼容、移动端测试、部署地址、运行说明", "完成毕业设计系统验收"),
            ("答辩交付", "演示视频、作品海报、最终PRD/设计文档", "用酸茶体验闭环讲清作品价值"),
        ],
        [1350, 4200, 3810],
    )

    doc.add_heading("10. 风险与对策", level=1)
    add_table(
        doc,
        ["风险", "影响", "对策"],
        [
            ("主题再次发散", "PRD变回泛茶文化或泛非遗平台", "所有一级模块以酸茶为主，史诗/民族文化只作为辅助解释"),
            ("H5游戏开发过重", "影响毕业设计完成度", "MVP优先完成技艺之魂、配方生成、海报和预约，其他场景可简化"),
            ("线下资源不可控", "核销体验难以实地验证", "使用可模拟核销流程与任务包原型，保留真实工坊合作为加分项"),
            ("电商范围膨胀", "购物、支付、售后过复杂影响主线", "只做酸茶相关商品与模拟支付，优先保证三类身份、商品、订单闭环"),
            ("权限设计混乱", "普通用户、商家、管理员数据互相越权", "从数据库和接口层区分role，后台菜单按角色加载，关键操作做权限校验"),
            ("文化表达失真", "影响项目可信度", "酸茶工艺、配方和话术需标注来源并经指导教师审核"),
            ("转化指标偏低", "商业闭环说服力不足", "优化海报吸引力、预约按钮位置、种子用户内测和KOL体验官招募"),
        ],
        [2200, 1800, 5360],
    )

    doc.add_heading("附录：重构原则", level=1)
    add_table(
        doc,
        ["原宽泛方向", "重构后处理"],
        [
            ("中国古代智慧体系/茶文化大主题", "收敛为德昂族酸茶制作技艺数字化再现"),
            ("完整茶文化展示", "仅保留与德昂酸茶来源、工艺、口味、体验有关的内容"),
            ("史诗独立叙事", "作为酸茶体验的引导、仪式和文化出处，不抢占主线"),
            ("普通文旅介绍站", "升级为H5轻游戏+预约核销+线下酸茶体验闭环"),
            ("单纯科普平台", "升级为科普内容、互动体验、预约核销、电商购买一体化平台"),
            ("毕业设计技术要求", "映射为15+页面、5+移动页、动态功能、数据库、兼容测试和发布"),
        ],
        [3300, 6060],
    )

    doc.save(OUT)


if __name__ == "__main__":
    build()
    print(OUT.resolve())
